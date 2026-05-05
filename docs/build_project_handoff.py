from docx import Document
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "docs/Sarath_Portfolio_Project_Handoff.docx"


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_code(doc, code):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    shade_cell(cell, "F3EFE6")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    for line_index, line in enumerate(code.splitlines()):
        if line_index:
            p.add_run().add_break()
        run = p.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(32, 31, 28)


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade_cell(cell, "EAF1F2")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(title + ": ")
    r.bold = True
    r.font.color.rgb = RGBColor(29, 93, 120)
    r.font.size = Pt(10)
    body_run = p.add_run(body)
    body_run.font.size = Pt(10)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.line_spacing = 1.08
styles["Normal"].paragraph_format.space_after = Pt(5)

for name, size, color in [
    ("Title", 25, RGBColor(18, 17, 15)),
    ("Heading 1", 16, RGBColor(18, 17, 15)),
    ("Heading 2", 12.5, RGBColor(201, 63, 45)),
]:
    style = styles[name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = True

title = doc.add_paragraph(style="Title")
title.add_run("Sarath Portfolio: Scroll Story Handoff")
subtitle = doc.add_paragraph()
subtitle_run = subtitle.add_run(
    "Project details and stickman animation guide for vibe coding"
)
subtitle_run.font.size = Pt(12)
subtitle_run.font.color.rgb = RGBColor(103, 97, 89)

meta = doc.add_table(rows=4, cols=2)
meta.columns[0].width = Inches(1.55)
meta.columns[1].width = Inches(5.7)
rows = [
    ("Project", "Minimal, non-conventional React portfolio that feels like a story."),
    ("Current URL", "http://localhost:5173/"),
    ("Stack", "React 19, TypeScript, Vite 5, CSS custom properties, scroll-driven animation."),
    ("Core files", "src/App.tsx, src/App.css, src/index.css, package.json."),
]
for row, values in zip(meta.rows, rows):
    set_cell_text(row.cells[0], values[0], True)
    set_cell_text(row.cells[1], values[1])
    shade_cell(row.cells[0], "F3EFE6")

doc.add_heading("1. Project Goal", level=1)
doc.add_paragraph(
    "The project is a minimalist, story-led portfolio for Sarath Krishnan PV. "
    "Instead of a conventional landing page with cards and sections, the page behaves like a scrolling narrative. "
    "A stickman narrator stays alive on screen while large text chapters move past him."
)
add_callout(
    doc,
    "Creative direction",
    "Keep the design sparse, personal, and strange in a controlled way. The content can evolve later; the first priority is motion, pacing, and story feeling.",
)

doc.add_heading("2. Stack and Why It Fits", level=1)
stack = doc.add_table(rows=1, cols=3)
headers = ["Part", "Used", "Reason"]
for i, h in enumerate(headers):
    set_cell_text(stack.rows[0].cells[i], h, True)
    shade_cell(stack.rows[0].cells[i], "12110F")
    stack.rows[0].cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 249, 238)
for part, used, reason in [
    ("App framework", "React", "Component state is useful for scroll position, movement phase, and future chapter choreography."),
    ("Language", "TypeScript", "Keeps animation state and future story data safer as the project grows."),
    ("Build tool", "Vite 5", "Fast dev server and simple static build for a portfolio."),
    ("Animation layer", "React + CSS variables", "No heavy animation dependency yet; scroll math happens in React and rendering happens in CSS."),
    ("3D", "Not used yet", "The current need is expressive 2D character motion. Three.js can be added later for depth or camera scenes."),
]:
    row = stack.add_row().cells
    set_cell_text(row[0], part)
    set_cell_text(row[1], used)
    set_cell_text(row[2], reason)

doc.add_heading("3. File Map", level=1)
files = doc.add_table(rows=1, cols=2)
for i, h in enumerate(["File", "Responsibility"]):
    set_cell_text(files.rows[0].cells[i], h, True)
    shade_cell(files.rows[0].cells[i], "F3EFE6")
for file_name, role in [
    ("src/App.tsx", "Story data, scroll listener, movement math, CSS variable output, page structure."),
    ("src/App.css", "Visual system, stickman drawing, limb transforms, diagonal travel, responsive layout."),
    ("src/index.css", "Global font, body reset, box sizing, selection color."),
    ("package.json", "Scripts and dependencies: dev, build, lint, preview."),
]:
    row = files.add_row().cells
    set_cell_text(row[0], file_name)
    set_cell_text(row[1], role)

doc.add_heading("4. Page Structure", level=1)
doc.add_paragraph(
    "The rendered page is built from two layers: a fixed stage layer and normal document sections."
)
add_bullets(
    doc,
    [
        "Fixed stage: the stickman, speech bubble, and scroll progress line. It stays on screen while the user scrolls.",
        "Story sections: intro, repeated chapter beats, and outro. Each section is very tall so scrolling becomes the main interaction.",
        "Oversized typography: the text is intentionally huge so the page feels like a kinetic story rather than a normal portfolio.",
    ],
)
add_code(
    doc,
    """<main className="story" style={cssVariables}>
  <aside className="stage">
    <div className="progress-track"><div className="progress-fill" /></div>
    <div className="stick-wrap">...</div>
  </aside>

  <section className="intro">...</section>
  {beats.map((beat) => <section className="beat">...</section>)}
  <section className="outro">...</section>
</main>""",
)

doc.add_heading("5. Stickman Animation: Full Explanation", level=1)
doc.add_heading("5.1 Scroll State", level=2)
doc.add_paragraph(
    "The animation is driven by the browser scroll position. React stores two values in state: progress and speed."
)
add_bullets(
    doc,
    [
        "progress: normalized scroll amount from 0 to 1.",
        "speed: approximate scroll velocity, clamped between -1 and 1. This makes the stickman lean when the user scrolls fast.",
        "requestAnimationFrame: used to batch scroll updates with the browser paint cycle and avoid excessive re-renders.",
    ],
)
add_code(
    doc,
    """const scrollable = document.documentElement.scrollHeight - window.innerHeight || 1
const speed = Math.max(-1, Math.min(distance / elapsed / 2, 1))

setMotion({
  progress: Math.min(y / scrollable, 1),
  speed,
})""",
)

doc.add_heading("5.2 Walking Cycle", level=2)
doc.add_paragraph(
    "The walking motion is a sine/cosine cycle. The scroll progress is converted into a repeating phase. "
    "That phase drives arms and legs in opposite directions."
)
add_code(
    doc,
    """const phase = motion.progress * Math.PI * 12
const walk = Math.sin(phase)
const counterWalk = Math.cos(phase)""",
)
doc.add_paragraph(
    "A sine wave creates smooth back-and-forth movement. The arms and legs use different wave values so the body does not move like a single flat object."
)

doc.add_heading("5.3 Limb Rotation", level=2)
doc.add_paragraph(
    "Each limb is a CSS element with transform-origin at the top. React calculates finished degree values and passes them to CSS variables."
)
add_code(
    doc,
    """const leftArm = -58 + counterWalk * 34
const rightArm = 58 + walk * 34
const leftLeg = -28 + walk * 36
const rightLeg = 28 + counterWalk * 36""",
)
add_bullets(
    doc,
    [
        "Base angle: gives each limb a natural resting direction.",
        "Wave multiplier: controls how big the swing is.",
        "Counter motion: prevents both arms or both legs from swinging together.",
    ],
)

doc.add_heading("5.4 Jumping", level=2)
doc.add_paragraph(
    "Jumping is not a separate animation timeline. It is another wave derived from scroll progress."
)
add_code(doc, "const jump = Math.max(0, Math.sin(motion.progress * Math.PI * 5))")
doc.add_paragraph(
    "The negative parts of the sine wave are removed with Math.max, so the value behaves like repeated jump arcs. "
    "CSS uses this to move the character upward and squash the shadow."
)

doc.add_heading("5.5 Lean From Scroll Speed", level=2)
add_code(doc, "const lean = motion.speed * 14")
doc.add_paragraph(
    "When the user scrolls quickly, speed becomes larger. The stickman rotates slightly so he feels like he is reacting to momentum."
)

doc.add_heading("5.6 Diagonal Movement", level=2)
doc.add_paragraph(
    "The character now travels diagonally across the viewport. React calculates a viewport X and Y from progress."
)
add_code(
    doc,
    """const diagonal = Math.sin(motion.progress * Math.PI * 2.4)
const diagonalLift = Math.cos(motion.progress * Math.PI * 2.4)
const travelX = 16 + ((diagonal + 1) / 2) * 58
const travelY = 34 + ((diagonalLift + 1) / 2) * 28""",
)
add_bullets(
    doc,
    [
        "travelX ranges roughly from 16vw to 74vw.",
        "travelY ranges roughly from 34vh to 62vh.",
        "Using both sine and cosine creates a smooth diagonal/orbital feel rather than simple left-to-right movement.",
    ],
)

doc.add_heading("5.7 Turning Around", level=2)
doc.add_paragraph(
    "The turn-around is created by flipping the stickman horizontally with scaleX. When the cosine changes sign, the character faces the other direction."
)
add_code(
    doc,
    """const face = Math.cos(motion.progress * Math.PI * 2.4) >= 0 ? 1 : -1

transform: translate(-50%, -50%)
  translateY(calc(var(--jump) * -34px))
  scaleX(var(--face))
  rotate(var(--lean));""",
)
add_callout(
    doc,
    "Important",
    "The flip is applied to the whole stickman, not each limb. This keeps the walking pose intact while changing the direction.",
)

doc.add_heading("5.8 CSS Figure Construction", level=2)
doc.add_paragraph(
    "The stickman is made from simple HTML spans. Each part is absolutely positioned inside .stickman."
)
add_bullets(
    doc,
    [
        ".head: circular outline with two pseudo-element eyes.",
        ".body: vertical rounded line.",
        ".arm and .leg: rounded bars rotating from the top.",
        ".shadow: ellipse under the feet, scaled during jumps.",
        ".speech: follows the same travel variables so it stays near the character.",
    ],
)

doc.add_heading("6. How To Choreograph Future Story Beats", level=1)
doc.add_paragraph(
    "The current beats array only contains text. The next useful evolution is to add an action field per story chapter."
)
add_code(
    doc,
    """const beats = [
  {
    marker: '01',
    action: 'jump-over-text',
    eyebrow: 'name enters',
    text: 'bro is sarath krishnan pv...'
  }
]""",
)
add_numbered(
    doc,
    [
        "Add action data to each beat.",
        "Calculate which beat is active based on scroll progress.",
        "Use the active action to change jump height, walking speed, face direction, z-index, text overlap, or special animations.",
        "Keep content data separate from motion math so story editing stays easy.",
    ],
)

doc.add_heading("7. Vibe Coding Prompt", level=1)
doc.add_paragraph(
    "Use this as a prompt when continuing the project with another coding agent:"
)
add_code(
    doc,
    """We are building a minimalist, non-conventional React portfolio for Sarath Krishnan PV.
The page should feel like a personal story, not a normal portfolio.

Current implementation:
- React + TypeScript + Vite.
- src/App.tsx owns story beats and scroll-driven animation math.
- src/App.css draws a stickman using HTML spans and CSS transforms.
- Scroll progress is normalized from 0 to 1.
- Walking is driven by sine/cosine phase math.
- The stickman receives CSS variables for arm angles, leg angles, jump height, lean, diagonal position, and facing direction.
- He moves diagonally across the screen and turns around with scaleX(var(--face)).

Creative direction:
- Keep typography huge and minimalist.
- Keep the stickman active, energetic, and story-aware.
- Add story chapters gradually.
- Choreograph motion per chapter: walk, jump over text, swing, curl, fall, run, pause, turn around.
- Avoid conventional portfolio cards unless there is a strong reason.""",
)

doc.add_heading("8. Commands", level=1)
commands = doc.add_table(rows=1, cols=2)
for i, h in enumerate(["Command", "Use"]):
    set_cell_text(commands.rows[0].cells[i], h, True)
    shade_cell(commands.rows[0].cells[i], "F3EFE6")
for command, use in [
    ("npm run dev", "Start local development server."),
    ("npm run build", "Type-check and create production build."),
    ("npm run lint", "Run ESLint checks."),
    ("npm run preview", "Preview production build locally."),
]:
    row = commands.add_row().cells
    set_cell_text(row[0], command)
    set_cell_text(row[1], use)

doc.add_heading("9. Current Known Constraints", level=1)
add_bullets(
    doc,
    [
        "The current stickman uses 2D CSS transforms. This is lightweight, but not physically realistic.",
        "The walking is scroll-progress driven, so it stops when scrolling stops.",
        "The jump timing is global, not tied to individual sections yet.",
        "The diagonal path can cross text. This is intentional, but future chapters may need section-specific z-index choreography.",
        "Three.js is not required yet. Add it later only if the story needs camera depth, parallax 3D scenes, or 3D objects.",
    ],
)

doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
doc.add_heading("Appendix: Key CSS Variable Contract", level=1)
contract = doc.add_table(rows=1, cols=3)
for i, h in enumerate(["Variable", "Source", "Effect"]):
    set_cell_text(contract.rows[0].cells[i], h, True)
    shade_cell(contract.rows[0].cells[i], "12110F")
    contract.rows[0].cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 249, 238)
for var, source, effect in [
    ("--progress", "scrollY / scrollable height", "Progress line height."),
    ("--jump", "positive sine wave", "Vertical hop and shadow squash."),
    ("--lean", "scroll speed", "Body tilt while scrolling fast."),
    ("--left-arm / --right-arm", "walk and counterWalk waves", "Arm swing."),
    ("--left-leg / --right-leg", "walk and counterWalk waves", "Leg swing."),
    ("--travel-x", "diagonal sine", "Horizontal viewport travel."),
    ("--travel-y", "diagonal cosine", "Vertical viewport travel."),
    ("--face", "cosine sign", "Turn-around / horizontal flip."),
]:
    row = contract.add_row().cells
    set_cell_text(row[0], var)
    set_cell_text(row[1], source)
    set_cell_text(row[2], effect)

doc.save(OUT)
print(OUT)
