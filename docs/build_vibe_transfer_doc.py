from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "docs/Sarath_Portfolio_Vibe_Coding_Transfer.docx"


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(9)
    if color:
        r.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def code(doc, text):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell, "F4F0E8")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    for i, line in enumerate(text.splitlines()):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        r.font.size = Pt(8.5)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.line_spacing = 1.08
styles["Normal"].paragraph_format.space_after = Pt(5)

for name, size, color in [
    ("Title", 24, RGBColor(18, 17, 15)),
    ("Heading 1", 16, RGBColor(18, 17, 15)),
    ("Heading 2", 12.5, RGBColor(185, 65, 46)),
]:
    style = styles[name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = color

doc.add_paragraph("Sarath Portfolio Transfer Doc", style="Title")
subtitle = doc.add_paragraph()
run = subtitle.add_run("Vibe-coding context to continue this project in another chat")
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(104, 96, 86)

summary = doc.add_table(rows=5, cols=2)
rows = [
    ("Current direction", "2D minimalist scroll-story portfolio with a CSS stickman narrator."),
    ("Do not continue", "The Three.js / cinematic / 3D attempts. User did not like them and asked to return to the doc checkpoint."),
    ("Visual tone", "Warm minimal editorial: cream paper background, black stickman, huge typography, playful micro-actions."),
    ("Current local URL", "http://localhost:5173/"),
    ("Main files", "src/App.tsx, src/App.css, src/index.css, docs/Sarath_Portfolio_Project_Handoff.docx"),
]
for row, values in zip(summary.rows, rows):
    set_text(row.cells[0], values[0], True)
    set_text(row.cells[1], values[1])
    shade(row.cells[0], "F4F0E8")

doc.add_heading("1. Current Project State", level=1)
doc.add_paragraph(
    "The app is a React + TypeScript + Vite portfolio prototype. The accepted direction is not a normal portfolio, "
    "not a 3D portfolio, and not a conventional landing page. It is a scroll-based life-story experience narrated by a stickman."
)
bullets(
    doc,
    [
        "Stickman is a fixed CSS character that travels while the user scrolls.",
        "Huge text sections tell Sarath's life in a playful narrative voice.",
        "The stickman's speech bubble adds sarcastic side commentary.",
        "The first page contains a slow, draggable, solar-system-like orbit of skill badges.",
        "The user prefers quirky but clean 2D animation over complex Three.js.",
    ],
)

doc.add_heading("2. What Was Tried and Rejected", level=1)
bullets(
    doc,
    [
        "Three.js cinematic portfolio with lighting, particles, post-processing, bazooka scene: rejected.",
        "Minimal black-and-white Three.js typography world: rejected.",
        "Moonwalk midway: removed because it caused glitches and broke flow.",
        "Bazooka in the 2D story: removed because simpler actions fit better.",
        "Public speaking skill badge/section: removed from current first-page skill system.",
    ],
)

doc.add_heading("3. Current Accepted Features", level=1)
features = doc.add_table(rows=1, cols=2)
set_text(features.rows[0].cells[0], "Feature", True, RGBColor(255, 250, 240))
set_text(features.rows[0].cells[1], "How it works", True, RGBColor(255, 250, 240))
shade(features.rows[0].cells[0], "12110F")
shade(features.rows[0].cells[1], "12110F")
for feature, detail in [
    ("Stickman narrator", "CSS spans for head/body/arms/legs. Position, limb angles, jump, lean, and action state are driven by CSS variables from React."),
    ("Scroll motion", "React listens to scroll, computes progress 0..1 and speed, then maps those to travelX/travelY, limb rotations, jump, and current beat."),
    ("Speech bubble", "Bubble text comes from the active beat's bubble field and follows the stickman."),
    ("Face expressions", "CSS changes mouth/eyes by data-action: wave, flowers, trip, point, lean."),
    ("Name section flowers", "The weird prop was removed. In the name section, small black line flowers pop around the stickman."),
    ("School section paper plane", "During action='trip', a paper plane crosses left to right, flips mid-flight, and leaves smoke circles."),
    ("First-page orbit", "Flutter, Coding, Video editing, and AI prompting badges orbit slowly like a simple solar system."),
    ("Orbit interactions", "User can drag the whole orbit to rotate it. User can drag each item to spin it individually."),
]:
    row = features.add_row().cells
    set_text(row[0], feature)
    set_text(row[1], detail)

doc.add_heading("4. First Page Skill Orbit", level=1)
doc.add_paragraph(
    "The first page includes all skill highlights. These should stay on the first page, not become separate story sections."
)
bullets(
    doc,
    [
        "Flutter: uses a Flutter-like blue mark based on the user's reference image.",
        "Coding: icon text is {/}.",
        "Video editing: simple video frame/play-style graphic.",
        "AI prompting: Gemini-style colorful sparkle icon based on the user's reference image.",
        "Orbit should remain slow. The user specifically said it was too fast when scrolling.",
        "Hover speed-up was removed because it caused glitches.",
    ],
)

doc.add_heading("5. Current Story Beats", level=1)
beats = doc.add_table(rows=1, cols=4)
for i, h in enumerate(["Marker", "Eyebrow", "Main Text", "Bubble / Action"]):
    set_text(beats.rows[0].cells[i], h, True, RGBColor(255, 250, 240))
    shade(beats.rows[0].cells[i], "12110F")
for marker, eyebrow, text, bubble_action in [
    ("00", "stick man speaks", "hi. i'm the stick man. walk with me. we're about to scroll through sarath's life, one strange little chapter at a time.", "yes, i am the narrator. / wave"),
    ("01", "name enters", "bro is sarath krishnan pv. born in 2001. a quiet start, a curious brain, and a suspicious amount of thinking.", "born in 2001. vintage enough. / flowers"),
    ("02", "school chapter", "then school happened. uniforms, corridors, notebooks, small wins, strange doubts, and the classic feeling that life forgot to give instructions.", "school: the original loading screen. / trip + paper plane"),
    ("03", "the turn", "somewhere between marks, machines, friends, and late nights, the story started choosing technology.", "ah yes. character development. / point"),
    ("04", "work in progress", "this portfolio is not a resume page. it is a walk. keep scrolling and the man keeps moving.", "resume pages could never. / lean"),
]:
    row = beats.add_row().cells
    set_text(row[0], marker)
    set_text(row[1], eyebrow)
    set_text(row[2], text)
    set_text(row[3], bubble_action)

doc.add_heading("6. Implementation Map", level=1)
bullets(
    doc,
    [
        "src/App.tsx owns story data, scroll state, current action, orbit drag state, and inline CSS variables.",
        "src/App.css owns all visual rendering: stickman, plane, flowers, orbit badges, huge typography, responsive layout.",
        "src/index.css contains global font/reset/background setup.",
        "No Three.js, GSAP, or Playwright dependencies are currently installed.",
        "package.json dependencies are only react and react-dom; Vite/TypeScript/ESLint are dev dependencies.",
    ],
)

doc.add_heading("7. Important Code Concepts", level=1)
doc.add_heading("Scroll Math", level=2)
code(
    doc,
    """progress = window.scrollY / (document height - viewport height)
speed = clamped scroll velocity
phase = progress * Math.PI * 12
walk = Math.sin(phase)
counterWalk = Math.cos(phase)

These values drive:
- arm/leg rotations
- jump arc
- lean
- travelX / travelY
- active beat and actionPulse""",
)
doc.add_heading("Action State", level=2)
code(
    doc,
    """currentBeatIndex = Math.floor(progress * beats.length)
currentAction = beats[currentBeatIndex].action
localProgress = progress * beats.length - currentBeatIndex
actionPulse = Math.sin(localProgress * Math.PI)

The main tag receives:
data-action={currentAction}

CSS selectors then style:
.story[data-action='flowers'] ...
.story[data-action='trip'] ...
.story[data-action='point'] ...""",
)
doc.add_heading("Orbit Drag", level=2)
doc.add_paragraph(
    "The orbit uses React state for manual rotation and item spins. Pointer events set dragRef and update either orbitRotation or itemSpins. "
    "CSS receives --orbit-rotation and --item-spin."
)

doc.add_heading("8. Next Best Tasks", level=1)
numbered(
    doc,
    [
        "Improve real content using Sarath's actual biography, schools, projects, and timeline.",
        "Polish first-page orbit positioning on mobile and desktop using screenshots.",
        "Make the stickman's motion more natural without returning to Three.js.",
        "Add one small action per story beat, but keep actions simple and clean.",
        "Replace placeholder email/contact details later.",
        "Update this transfer doc after major changes.",
    ],
)

doc.add_heading("9. Prompt for the Next Chat", level=1)
code(
    doc,
    """We are continuing a React + TypeScript + Vite portfolio in D:\\D\\portfolio.
Do not rebuild it in Three.js. The user rejected the 3D options.

Current accepted direction:
- 2D minimalist scroll-story portfolio.
- Cream paper background, huge typography, black CSS stickman narrator.
- Stickman walks while scrolling and gives sarcastic bubble commentary.
- First page has slow draggable solar-system skill orbit:
  Flutter, Coding ({/}), Video editing, AI prompting/Gemini sparkle.
- User can drag the orbit and spin individual items.
- School section has paper plane crossing left to right with smoke trail.
- Name section has flowers popping. No weird gray prop.
- No bazooka. No moonwalk. No public speaking skill.

Before changing anything, inspect src/App.tsx and src/App.css.
Keep changes small, preserve the current vibe, and verify with npm run lint and npm run build.""",
)

doc.add_heading("10. Verification Commands", level=1)
code(
    doc,
    """npm run dev
npm run lint
npm run build""",
)

doc.save(OUT)
print(OUT)
